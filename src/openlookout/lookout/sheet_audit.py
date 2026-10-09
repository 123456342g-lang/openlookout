"""手工汇总核对：一个工作簿里有好几张同样表头的明细表（每周一张），再加一张手打的“汇总”表（没有公式）。
检测：按明细表重新加一遍，和手打的汇总逐格比；能用一条简单折算规则（比如“每 k 次 A 折成 1 次 B”）解释的差异单独标出来。
求解：生成一份新工作簿（原表原样保留 + 一张用公式自动汇总的表 + 一张差异核对表），用 LibreOffice 无界面重算公式，
      和 Python 独立重算的结果逐格比对。不修改用户的原文件。纯规则，不用模型。"""
import hashlib, re, shutil, subprocess, tempfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter, quote_sheetname

SUMMARY_NAME = re.compile(r"汇总|合计|总表|统计|summary|total", re.I)

def _num(v):
    if v in (None, ""): return 0.0
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).strip())
    except ValueError: return None

def _header(ws):
    for i, r in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True), 1):
        cells = [str(c).strip() if c is not None else "" for c in r]
        if sum(1 for c in cells if c and not re.fullmatch(r"[\d.]+", c)) >= 3:
            while cells and not cells[-1]: cells.pop()
            return i, cells
    return None, None

def _rows(ws, hrow, hdr, key_i, num_cols):
    out = []
    for r in ws.iter_rows(min_row=hrow + 1, values_only=True):
        k = str(r[key_i]).strip() if key_i < len(r) and r[key_i] not in (None, "") else ""
        if not k: continue
        vals = {c: _num(r[i]) if i < len(r) else 0.0 for c, i in num_cols.items()}
        if all(v in (0.0, None) for v in vals.values()): continue    # “全勤”这类说明行
        out.append((k, vals, r))
    return out

def t_is_text(col): return bool(re.search(r"备注|说明|原因|日期|时间|姓名|名称|地点|note|remark|date", col, re.I))

def analyse(path):
    """返回 None（不是这种结构）或分析结果 dict。"""
    try: wb = openpyxl.load_workbook(path)          # 公式要看原样，不能 data_only
    except Exception: return None
    sheets = {}
    for ws in wb.worksheets:
        h, hdr = _header(ws)
        if h: sheets[ws.title] = (ws, h, hdr)
    summ = [t for t in sheets if SUMMARY_NAME.search(t)]
    if not summ: return None
    st = summ[0]; sws, sh, shdr = sheets[st]
    details = [t for t, (_, _, hdr) in sheets.items() if t != st and hdr[:len(shdr)] == shdr[:len(hdr)] and len(set(hdr) & set(shdr)) >= 3]
    if len(details) < 3: return None
    key_i = 0; key = shdr[0]
    # 数值列：明细里出现过数值，且汇总里也有的列
    num_cols = {}
    for i, c in enumerate(shdr):
        if i == key_i or not c: continue
        vals = [r[i] for t in details + [st] for r in sheets[t][0].iter_rows(min_row=sheets[t][1] + 1, values_only=True) if i < len(r) and r[i] not in (None, "")]
        if t_is_text(c): continue
        if vals and sum(_num(v) is not None for v in vals) >= 0.9 * len(vals): num_cols[c] = i
    if len(num_cols) < 2: return None
    formulas = sum(1 for r in sws.iter_rows(min_row=sh + 1) for c in r if isinstance(c.value, str) and c.value.startswith("="))
    # 明细重算
    calc = defaultdict(lambda: defaultdict(float)); contrib = defaultdict(list)
    for t in details:
        ws, h, hdr = sheets[t]
        for k, vals, raw in _rows(ws, h, hdr, key_i, num_cols):
            for c, v in vals.items():
                if v: calc[k][c] += v; contrib[(k, c)].append((t, v))
    hand = {}
    for k, vals, raw in _rows(sws, sh, shdr, key_i, num_cols): hand[k] = vals
    cols = list(num_cols)
    def diffs(table):
        return [(k, c, table.get(k, {}).get(c, 0.0), hand.get(k, {}).get(c, 0.0) or 0.0)
                for k in sorted(set(table) | set(hand)) for c in cols if abs(table.get(k, {}).get(c, 0.0) - (hand.get(k, {}).get(c, 0.0) or 0.0)) > 1e-9]
    raw_diffs = diffs(calc)
    # 能不能用“每 k 次 A 折成 1 次 B”解释？只在它解释了足够多行、且不弄坏原本对得上的行时才采用
    rule = None
    if raw_diffs:
        bad_rows = {k for k, *_ in raw_diffs}
        best = None
        for a in cols:
            for b in cols:
                if a == b: continue
                for kk in range(2, 6):
                    adj = {n: dict(v) for n, v in calc.items()}
                    for n, v in adj.items():
                        q = int(v.get(a, 0) // kk)
                        if q: v[a] = v.get(a, 0) - q * kk; v[b] = v.get(b, 0) + q
                    left = diffs(adj); left_rows = {k for k, *_ in left}
                    fixed = len(bad_rows - left_rows); broke = len(left_rows - bad_rows)
                    if fixed >= 3 and broke == 0 and fixed >= 0.6 * len(bad_rows) and (best is None or fixed > best[0]):
                        best = (fixed, a, b, kk, adj, left)
        if best:
            fixed, a, b, kk, adj, left = best
            rule = {"from": a, "to": b, "k": kk, "rows_explained": fixed, "rows_with_diff": len(bad_rows), "left": left, "adjusted": adj}
    return {"path": str(path), "summary_sheet": st, "summary_header_row": sh, "header": shdr, "details": details, "key": key,
            "cols": cols, "num_cols": num_cols, "summary_formulas": formulas, "calc": {k: dict(v) for k, v in calc.items()}, "hand": hand,
            "raw_diffs": raw_diffs, "rule": rule, "contrib": {f"{k}|{c}": v for (k, c), v in contrib.items()},
            "sha1": hashlib.sha1(Path(path).read_bytes()).hexdigest()}

def detect_manual_summary(evs, meta):
    from .common import watched_files, ROOT
    home = ROOT / meta["user_home"]
    seen, out = {}, []
    last = {}
    for e in evs:
        if e.get("path") and e["type"] in ("file_save", "file_create"): last[e["path"]] = max(last.get(e["path"], e["dt"]), e["dt"])
    for p in sorted(watched_files(meta)):
        if p.suffix.lower() != ".xlsx": continue
        A = analyse(p)
        if not A or A["summary_formulas"]: continue
        rel = p.relative_to(home).as_posix()
        if A["sha1"] in seen: seen[A["sha1"]]["local"]["copies"].append(rel); continue
        rows_diff = len({k for k, *_ in A["raw_diffs"]})
        left = A["rule"]["left"] if A["rule"] else A["raw_diffs"]
        rows_left = len({k for k, *_ in left})
        f = {"kind": "manual_summary", "score": 15 + 4 * min(rows_left, 8) + min(rows_diff, 10),
             "local": {"file": rel, "copies": [rel], "analysis": A, "last_saved": str(last.get(rel, ""))[:16]},
             "facts": {"detail_sheets": len(A["details"]), "people_rows": len(set(A["calc"]) | set(A["hand"])), "value_columns": len(A["cols"]),
                       "summary_has_formulas": False, "cells_differ": len(A["raw_diffs"]), "rows_differ": rows_diff,
                       "rule_found": bool(A["rule"]), "rule_desc": (f"每 {A['rule']['k']} 次「列{A['cols'].index(A['rule']['from'])+1}」折成 1 次「列{A['cols'].index(A['rule']['to'])+1}」") if A["rule"] else None,
                       "rows_explained_by_rule": A["rule"]["rows_explained"] if A["rule"] else 0, "rows_unexplained": rows_left,
                       "cells_unexplained": len(left), "identical_copies": 1}}
        seen[A["sha1"]] = f; out.append(f)
    for f in out: f["facts"]["identical_copies"] = len(f["local"]["copies"])
    out = [f for f in out if f["facts"]["cells_differ"]]
    # 同样结构的几个工作簿（比如第一、第二学期的同一种表）是同一件事，合成一条，免得一件事上两张卡
    groups = defaultdict(list)
    for f in out:
        A = f["local"]["analysis"]; groups[(tuple(A["header"]), A["summary_sheet"], str(Path(f["local"]["file"]).parent))].append(f)
    merged = []
    for fs in groups.values():
        if len(fs) == 1: fs[0]["local"]["parts"] = [dict(fs[0]["local"])]; merged.append(fs[0]); continue
        F = {k: sum(f["facts"][k] for f in fs) for k in ("detail_sheets", "people_rows", "cells_differ", "rows_differ", "rows_explained_by_rule", "rows_unexplained", "cells_unexplained")}
        F.update({"value_columns": max(f["facts"]["value_columns"] for f in fs), "summary_has_formulas": False, "workbooks": len(fs),
                  "rule_found": any(f["facts"]["rule_found"] for f in fs), "rule_desc": next((f["facts"]["rule_desc"] for f in fs if f["facts"]["rule_desc"]), None),
                  "rule_workbooks": sum(1 for f in fs if f["facts"]["rule_found"]), "identical_copies": min(f["facts"]["identical_copies"] for f in fs)})
        merged.append({"kind": "manual_summary", "score": max(f["score"] for f in fs) + 5 * (len(fs) - 1),
                       "local": {"file": fs[0]["local"]["file"], "copies": [c for f in fs for c in f["local"]["copies"]], "parts": [f["local"] for f in fs],
                                 "analysis": fs[0]["local"]["analysis"], "last_saved": max(f["local"]["last_saved"] for f in fs)}, "facts": F})
    for f in merged: f["facts"].setdefault("workbooks", 1); f["facts"].setdefault("rule_workbooks", int(f["facts"]["rule_found"]))
    return merged

# ---------------------------------------------------------------- 求解
def build(A, out_xlsx, rule=True):
    """原表原样复制 + 「自动汇总(公式)」 + 「差异核对」。返回写入的期望值表（供验证）。"""
    src = Path(A["path"]); shutil.copyfile(src, out_xlsx)
    wb = openpyxl.load_workbook(out_xlsx)
    cols, num_cols, details = A["cols"], A["num_cols"], A["details"]
    keys = list(A["hand"]) + [k for k in A["calc"] if k not in A["hand"]]
    ws = wb.create_sheet("自动汇总(公式)")
    ws.append(["本表由公式从各周明细自动加总，明细改了这里会跟着变（OpenLookout 生成，原表未改动）"])
    ws.append([A["key"]] + cols)
    expect = {}
    for r, k in enumerate(keys, start=3):
        ws.cell(r, 1, k)
        for j, c in enumerate(cols, start=2):
            letter = get_column_letter(num_cols[c] + 1)
            parts = [f"SUMIF({quote_sheetname(t)}!$A:$A,$A{r},{quote_sheetname(t)}!{letter}:{letter})" for t in details]
            ws.cell(r, j, "=" + "+".join(parts))
            expect[(k, c)] = A["calc"].get(k, {}).get(c, 0.0)
    R = A["rule"] if rule else None
    if R:
        ws2 = wb.create_sheet("按折算规则(公式)")
        ws2.append([f"推测的折算规则：每 {R['k']} 次「{R['from']}」记 1 次「{R['to']}」（从你原来的汇总里反推出来的，不确定是否是正式规定）"])
        ws2.append([A["key"]] + cols)
        src_name = quote_sheetname("自动汇总(公式)")
        ci = {c: get_column_letter(j) for j, c in enumerate(cols, start=2)}
        for r, k in enumerate(keys, start=3):
            ws2.cell(r, 1, k)
            for c in cols:
                base = f"{src_name}!{ci[c]}{r}"
                if c == R["from"]: fml = f"=MOD({base},{R['k']})"
                elif c == R["to"]: fml = f"={base}+INT({src_name}!{ci[R['from']]}{r}/{R['k']})"
                else: fml = f"={base}"
                ws2.cell(r, cols.index(c) + 2, fml)
                expect[("rule", k, c)] = R["adjusted"].get(k, {}).get(c, 0.0)
    ws3 = wb.create_sheet("差异核对")
    ws3.append([A["key"], "项目", "明细加总", "原汇总手填", "差值", "能否用折算规则解释", "明细来源（工作表:数值）"])
    left = {(k, c) for k, c, *_ in (R["left"] if R else A["raw_diffs"])}
    for k, c, calc_v, hand_v in A["raw_diffs"]:
        ws3.append([k, c, calc_v, hand_v, hand_v - calc_v, ("否，需人工核对" if (k, c) in left else "是") if R else "（未发现规则）",
                    "；".join(f"{t}:{v:g}" for t, v in A["contrib"].get(f"{k}|{c}", [])) or "明细里没有记录"])
    wb.save(out_xlsx)
    return expect, keys

def recalc_and_read(xlsx, workdir):
    """LibreOffice 无界面打开 → 重算 → 另存，读回公式结果。"""
    workdir = Path(workdir); workdir.mkdir(parents=True, exist_ok=True)
    prof = Path(tempfile.mkdtemp(prefix="lo_prof_"))
    cmd = ["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(workdir), str(xlsx)]
    subprocess.run(cmd, check=True, capture_output=True, timeout=180)
    shutil.rmtree(prof, ignore_errors=True)
    return openpyxl.load_workbook(workdir / Path(xlsx).name, data_only=True)

def verify(A, out_xlsx, expect, keys, workdir):
    res = {"cells_checked": 0, "cells_match": 0, "rule_cells_checked": 0, "rule_cells_match": 0, "original_sheets_unchanged": None, "recalc_engine": "LibreOffice headless"}
    try: wb = recalc_and_read(out_xlsx, workdir)
    except Exception as ex:
        res["recalc_engine"] = f"不可用：{type(ex).__name__}"; return res
    ws = wb["自动汇总(公式)"]
    for r, k in enumerate(keys, start=3):
        assert ws.cell(r, 1).value == k
        for j, c in enumerate(A["cols"], start=2):
            res["cells_checked"] += 1; v = ws.cell(r, j).value or 0
            if abs(float(v) - expect[(k, c)]) < 1e-9: res["cells_match"] += 1
    if "按折算规则(公式)" in wb.sheetnames:
        ws = wb["按折算规则(公式)"]
        for r, k in enumerate(keys, start=3):
            for j, c in enumerate(A["cols"], start=2):
                res["rule_cells_checked"] += 1; v = ws.cell(r, j).value or 0
                if abs(float(v) - expect[("rule", k, c)]) < 1e-9: res["rule_cells_match"] += 1
    # 原有工作表逐格不变
    a = openpyxl.load_workbook(A["path"]); b = openpyxl.load_workbook(out_xlsx)
    same = all([[c for c in r] for r in a[t].iter_rows(values_only=True)] == [[c for c in r] for r in b[t].iter_rows(values_only=True)] for t in a.sheetnames)
    res["original_sheets_unchanged"] = same
    return res
