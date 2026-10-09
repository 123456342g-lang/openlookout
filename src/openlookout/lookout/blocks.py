"""积木库：每块都是我们写好、测过、行为确定的小函数。云端模型只能"拼"这些积木。
权限声明：所有积木只读输入、只在输出目录写新文件；没有一块会发送、提交、付款或删除。"""
import builtins, json, re, statistics, subprocess, tempfile
from datetime import datetime
from pathlib import Path
import openpyxl
from .common import read_table, to_num, to_date

BLOCK_DOCS = {
 "read_table": "read_table(source:'表X') -> 表格(行列表)。读 xlsx/csv。",
 "filter_date_range": "filter_date_range(table, column, start, end) -> 表格。保留 column 日期在 [start, end] 内的行（含两端）。start/end 可用 {week_start} {week_end} {today}。",
 "filter_rows": "filter_rows(table, column, op, value) -> 表格。op ∈ eq, ne, contains, not_contains, in。",
 "aggregate": "aggregate(table, op, column=None, round=None) -> 数值。op ∈ count(行数), nunique(不同值个数), sum, mean, min, max。",
 "select_columns": "select_columns(table, columns:[...], rename:{旧名:新名}=None) -> 表格。",
 "sort_rows": "sort_rows(table, column, desc=false) -> 表格。",
 "lookup_map": "lookup_map(table, key_column, value_column) -> 映射。从表格里建立 键→值 对照（同键多值会记为冲突）。",
 "extract_fields": "extract_fields(docs:'文档组X', fields:{字段名: 正则(恰好一个捕获组)}) -> 表格。每个文档一行，另有 _doc 列。",
 "map_to_table": "map_to_table(map, key_name, value_name) -> 表格。把 lookup_map 的结果变成两列表格。",
 "write_xlsx_report": "write_xlsx_report(filename, title, cells:[{label, value}], table:{ref:'$步骤id(必须是表格)', columns:[...]}=None, period=None) -> 文件。第1行标题；period 写成 A2'统计周期'/B2；cells 从第4行起 A列写 label、B列写 value；table 写在下方。",
 "write_xlsx_table": "write_xlsx_table(table:'$步骤id(必须是表格)', filename, title=None) -> 文件。",
 "make_form_fill_script": "make_form_fill_script(table, field_map:{网页字段标签: 表格列名}, filename) -> 文件。生成浏览器脚本：按字段标签把一行数据填进网页表单；永远不点提交。",
 "python_transform": "python_transform(code, inputs:{名字: $步骤}) -> 任意。code 里定义 def transform(**inputs)，只能做数据变换；在无网络沙箱里单独运行，超时 10 秒。",
}

class Ctx:
    def __init__(self, refs, params, outdir):
        self.refs = refs; self.params = params; self.out = Path(outdir); self.vals = {}; self.files = []

def _resolve(ctx, v):
    if isinstance(v, str):
        if v.startswith("$") and v[1:] in ctx.vals: return ctx.vals[v[1:]]
        for k, p in ctx.params.items(): v = v.replace("{" + k + "}", str(p))
        return v
    if isinstance(v, list): return [_resolve(ctx, x) for x in v]
    if isinstance(v, dict): return {k: _resolve(ctx, x) for k, x in v.items()}
    return v

def _norm_cell(v):
    if isinstance(v, datetime): return v.date().isoformat()
    return v

def b_read_table(ctx, source):
    h, rows = read_table(ctx.refs[source])
    return [{h[i]: _norm_cell(r[i]) if i < len(r) else None for i in range(len(h))} for r in rows]

def b_filter_date_range(ctx, table, column, start, end):
    s, e = to_date(start), to_date(end)
    return [r for r in table if (d := to_date(r.get(column))) and s <= d <= e]

def b_filter_rows(ctx, table, column, op, value):
    f = {"eq": lambda x: str(x) == str(value), "ne": lambda x: str(x) != str(value),
         "contains": lambda x: str(value) in str(x or ""), "not_contains": lambda x: str(value) not in str(x or ""),
         "in": lambda x: str(x) in [str(v) for v in value]}[op]
    return [r for r in table if f(r.get(column))]

def b_aggregate(ctx, table, op, column=None, round=None):
    if op == "count": v = len(table)
    elif op == "nunique": v = len({r.get(column) for r in table})
    else:
        xs = [to_num(r.get(column)) for r in table if to_num(r.get(column)) is not None]
        v = {"sum": sum, "mean": lambda a: statistics.mean(a) if a else 0, "min": min, "max": max}[op](xs) if xs or op in ("sum", "mean") else None
    if round is not None and v is not None: v = builtins.round(v, int(round))
    if isinstance(v, float) and v.is_integer() and op in ("sum", "count", "nunique"): v = int(v)
    return v

def b_select_columns(ctx, table, columns, rename=None):
    rename = rename or {}
    return [{rename.get(c, c): r.get(c) for c in columns} for r in table]

def b_sort_rows(ctx, table, column, desc=False):
    return sorted(table, key=lambda r: str(r.get(column)), reverse=bool(desc))

def b_lookup_map(ctx, table, key_column, value_column):
    m, conflicts = {}, {}
    for r in table:
        k, v = r.get(key_column), r.get(value_column)
        if k in (None, "") or v in (None, ""): continue
        if k in m and m[k] != v: conflicts.setdefault(k, {m[k]}).add(v)
        m.setdefault(k, v)
    return {"map": m, "conflicts": {k: sorted(map(str, v)) for k, v in conflicts.items()}}

def b_extract_fields(ctx, docs, fields):
    rows = []
    for p in sorted(Path(x) for x in ctx.refs[docs]):
        t = p.read_text(encoding="utf-8", errors="ignore"); row = {"_doc": p.name}
        for k, rx in fields.items():
            m = re.search(rx, t); row[k] = m.group(1).strip() if m else None
        rows.append(row)
    return rows

def b_write_xlsx_report(ctx, filename, title, cells, table=None, period=None):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "周报"
    ws["A1"] = title
    if period: ws["A2"] = "统计周期"; ws["B2"] = period
    r = 4
    for c in cells: ws.cell(r, 1, c["label"]); ws.cell(r, 2, c["value"]); r += 1
    if table:
        r += 1; cols = table["columns"]; rows = _as_table(table["ref"])
        if not isinstance(rows, list) or (rows and not isinstance(rows[0], dict)):
            raise ValueError("write_xlsx_report 的 table.ref 必须引用一个表格步骤（如 \"$s5\"）")
        for j, c in enumerate(cols): ws.cell(r, 1 + j, c)
        for i, row in enumerate(rows):
            for j, c in enumerate(cols): ws.cell(r + 1 + i, 1 + j, row.get(c))
    p = ctx.out / filename; wb.save(p); ctx.files.append(str(p)); return str(p)

def _as_table(t):
    """迭代 2：积木对常见的"差一点对"的输入更宽容（模型常把 {columns, rows} 或映射直接交过来）。"""
    if isinstance(t, dict) and "rows" in t and "columns" in t:
        return [dict(zip(t["columns"], r)) if isinstance(r, (list, tuple)) else r for r in t["rows"]]
    if isinstance(t, dict) and "map" in t: t = t["map"]
    if isinstance(t, dict): return [{"键": k, "值": v} for k, v in t.items()]
    return t

def b_map_to_table(ctx, map, key_name, value_name):
    m = map.get("map", map) if isinstance(map, dict) else map
    return [{key_name: k, value_name: v} for k, v in m.items()]

def b_write_xlsx_table(ctx, table, filename, title=None):
    table = _as_table(table)
    wb = openpyxl.Workbook(); ws = wb.active; r = 1
    if title: ws.cell(1, 1, title); r = 3
    cols = list(table[0].keys()) if table else []
    for j, c in enumerate(cols): ws.cell(r, 1 + j, c)
    for i, row in enumerate(table):
        for j, c in enumerate(cols): ws.cell(r + 1 + i, 1 + j, row.get(c))
    p = ctx.out / filename; wb.save(p); ctx.files.append(str(p)); return str(p)

FILL_JS = r"""// OpenLookout 生成的填表脚本：只填写，不提交。提交按钮永远由你自己按。
(function(){
  const DATA = __DATA__;            // 每张待报发票一行
  const MAP = __MAP__;              // 网页字段标签 -> 数据列
  const idx = Number(new URLSearchParams(location.search).get('ol_row') || 0);
  const row = DATA[idx]; if(!row) return;
  const labels = Array.from(document.querySelectorAll('label'));
  let filled = 0;
  for (const [label, col] of Object.entries(MAP)) {
    const lb = labels.find(l => l.textContent.replace(/[\s：:*]/g,'') === label);
    const input = lb && (lb.htmlFor ? document.getElementById(lb.htmlFor) : lb.querySelector('input'));
    if (!input) continue;
    input.value = row[col] ?? ''; input.dispatchEvent(new Event('input', {bubbles:true})); filled++;
  }
  document.body.setAttribute('data-ol-filled', String(filled));
})();
"""
def b_make_form_fill_script(ctx, table, field_map, filename):
    js = FILL_JS.replace("__DATA__", json.dumps(table, ensure_ascii=False)).replace("__MAP__", json.dumps(field_map, ensure_ascii=False))
    p = ctx.out / filename; p.write_text(js, encoding="utf-8"); ctx.files.append(str(p)); return str(p)

def b_python_transform(ctx, code, inputs):
    """模型写的代码：单独进程 + bubblewrap（无网络、只读文件系统、只有一个临时目录可写）+ 10 秒超时。"""
    with tempfile.TemporaryDirectory() as td:
        Path(td, "code.py").write_text(code + "\nimport json,sys\nd=json.load(open(sys.argv[1]))\nprint(json.dumps(transform(**d),ensure_ascii=False,default=str))\n")
        Path(td, "in.json").write_text(json.dumps(inputs, ensure_ascii=False, default=str))
        cmd = ["bwrap", "--ro-bind", "/usr", "/usr", "--ro-bind", "/lib", "/lib", "--ro-bind-try", "/lib64", "/lib64",
               "--ro-bind", "/bin", "/bin", "--ro-bind", td, "/job", "--tmpfs", "/tmp", "--proc", "/proc", "--dev", "/dev",
               "--unshare-all", "--die-with-parent", "--clearenv", "python3", "-I", "/job/code.py", "/job/in.json"]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=10)
        if r.returncode: raise RuntimeError("python_transform 失败: " + r.stderr[-300:])
        return json.loads(r.stdout)

BLOCKS = {k: globals()["b_" + k] for k in BLOCK_DOCS}

TABLE_ARGS = {"table", "map"}
def validate(wf):
    import inspect
    errs = []
    ids = set()
    for s in wf.get("steps", []):
        if s.get("block") not in BLOCKS: errs.append(f"未知积木 {s.get('block')}"); continue
        params = inspect.signature(BLOCKS[s["block"]]).parameters
        for a, v in (s.get("args") or {}).items():
            if a not in params: errs.append(f"{s.get('id')}: 积木 {s['block']} 没有参数 {a}")
            if a in TABLE_ARGS and s["block"] != "write_xlsx_report" and not (isinstance(v, str) and v.startswith("$")): errs.append(f"{s.get('id')}: 参数 {a} 必须是 '$步骤id'")
        tb = (s.get("args") or {}).get("table")
        if s["block"] == "write_xlsx_report" and isinstance(tb, dict) and not str(tb.get("ref", "")).startswith("$"):
            errs.append(f"{s.get('id')}: table.ref 必须是 '$步骤id'，不能是单元格地址")
        ids.add(s.get("id"))
        for v in json.dumps(s.get("args", {}), ensure_ascii=False).split('"'):
            if v.startswith("$") and v[1:] not in ids: errs.append(f"{s.get('id')} 引用了不存在/靠后的步骤 {v}")
    return errs

def run_workflow(wf, refs, params, outdir):
    ctx = Ctx(refs, params, outdir); Path(outdir).mkdir(parents=True, exist_ok=True); trace = {}
    for s in wf["steps"]:
        args = _resolve(ctx, s.get("args", {}))
        ctx.vals[s["id"]] = BLOCKS[s["block"]](ctx, **args)
        trace[s["id"]] = {"block": s["block"], "args": s.get("args", {})}
    return ctx.vals, ctx.files, trace
