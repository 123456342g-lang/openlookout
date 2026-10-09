"""采集层输出的读取 + 通用表格读取。只读 capture/ 与用户磁盘上的文件，不读 ground_truth/。"""
import csv, io, json, re
from datetime import datetime
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
import os
CAP = Path(os.environ.get("LOOKOUT_CAPTURE", ROOT / "capture"))  # 可指向真实记录器的输出目录
SIM = (ROOT / os.environ["LOOKOUT_HOME"]).parent if os.environ.get("LOOKOUT_HOME") else ROOT / "sim_home"

def load_events():
    evs = []
    for line in open(CAP / "events.jsonl", encoding="utf-8"):
        if not line.strip(): continue
        e = json.loads(line); e["dt"] = datetime.fromisoformat(e["ts"]).replace(tzinfo=None); evs.append(e)
    evs.sort(key=lambda e: e["dt"])   # 真实记录器在"离开窗口"时才写 focus 事件，必须按时间重排
    return evs

def load_meta():
    mp = CAP / "meta.json"
    if mp.exists(): m = json.loads(mp.read_text(encoding="utf-8"))
    else:  # 真实记录：没有 meta 就用"现在"和默认监听目录
        m = {"now": datetime.now().replace(microsecond=0).isoformat(), "user_home": os.environ.get("LOOKOUT_HOME", "sim_home/xiaoyu"), "watched_dirs": ["文档", "下载"]}
    m["now_dt"] = datetime.fromisoformat(m["now"])
    # 用户在设置里填的自己的称呼（用于判断"这条是不是派给你的"）
    m.setdefault("user_aliases", ["林晓雨", "晓雨"])
    return m

def abspath(rel):
    return SIM / rel

# 应用缓存目录里的文件不是用户写的（真实 Windows 回溯里出现过聊天软件内嵌浏览器的缓存 index.txt）
CACHE_DIR = re.compile(r"(/|^)(AppData|Service Worker|CacheStorage|Cache|qtCef|GPUCache|Code Cache|node_modules|\.git)(/|$)", re.I)
def watched_files(meta):
    home = ROOT / meta["user_home"]
    out = []
    for d in meta["watched_dirs"]:
        out += [p for p in (home / d).rglob("*") if p.is_file() and not CACHE_DIR.search(p.relative_to(home).as_posix())]
    return out

def read_text_any(path):
    """读文本类文件：txt/md 依次试 utf-8 / utf-16(带 BOM) / gb18030（Windows 记事本常见）；docx 取段落和表格文字。读不出来返回空串。"""
    path = Path(path)
    if path.suffix.lower() == ".docx":
        try:
            import docx
            d = docx.Document(str(path))
            lines = [p.text for p in d.paragraphs]
            for t in d.tables:
                for r in t.rows: lines.append(" | ".join(c.text for c in r.cells))
            return "\n".join(lines)
        except Exception:
            import zipfile
            try:
                xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8", "ignore")
                return re.sub(r"<[^>]+>", "", xml.replace("</w:p>", "\n"))
            except Exception: return ""
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")): return raw.decode("utf-16", "replace")
    for enc in ("utf-8-sig", "gb18030"):
        try: return raw.decode(enc)
        except UnicodeDecodeError: continue
    return raw.decode("utf-8", "replace")

def norm_title(t):
    """窗口标题归一：数字串→#，便于把"发票_123.pdf"与"发票_456.pdf"视为同一类窗口"""
    return re.sub(r"\d+", "#", t or "")

def read_table(path):
    """读 xlsx/csv 为 (header, rows)；自动跳过账单类文件前面的说明行。"""
    path = Path(path)
    if path.suffix == ".xlsx":
        ws = openpyxl.load_workbook(path, data_only=True).active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "gbk"):
            try: txt = raw.decode(enc); break
            except UnicodeDecodeError: continue
        rows = list(csv.reader(io.StringIO(txt)))
    # 表头 = 第一个"之后连续若干行列数与它相同"的行
    best = 0
    for i, r in enumerate(rows[:30]):
        n = sum(1 for c in r if c not in (None, ""))
        if n >= 3 and all(len([c for c in rr if c not in (None, "")]) >= n - 1 for rr in rows[i + 1:i + 4]) and len(rows) > i + 1:
            best = i; break
    header = [str(c).strip() if c is not None else "" for c in rows[best]]
    data = [r for r in rows[best + 1:] if any(c not in (None, "") for c in r)]
    return header, data

def to_num(v):
    if isinstance(v, (int, float)): return float(v)
    if v is None: return None
    s = str(v).replace(",", "").replace("¥", "").strip()
    try: return float(s)
    except ValueError: return None

DATE_RE = re.compile(r"(20\d\d)[-/.年](\d{1,2})[-/.月](\d{1,2})")
def to_date(v):
    if isinstance(v, datetime): return v
    m = DATE_RE.search(str(v or ""))
    if not m: return None
    try: return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError: return None
